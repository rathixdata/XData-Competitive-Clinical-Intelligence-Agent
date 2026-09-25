"""Exports with source references and generation timestamp (FR-QA-006, FR-ALT-003).

DOCX (python-docx), PDF (fpdf2), PowerPoint-ready PPTX (python-pptx) and Markdown.
Statement types are always labelled so exported intelligence keeps fact/inference/unknown separation.
"""

from __future__ import annotations

import io
from typing import Any

TYPE_LABEL = {"FACT": "Verified fact", "INFERENCE": "AI interpretation", "UNKNOWN": "Unknown",
              "RECOMMENDED_INVESTIGATION": "Recommended investigation"}


def _doc_model(kind: str, content: dict[str, Any]) -> dict[str, Any]:
    """Normalize answers / narratives / briefs to a common export structure."""
    if kind == "ask_answer":
        return {
            "title": f"Ask-the-Landscape: {content.get('question', '')}",
            "summary": content.get("answer", ""),
            "statements": content.get("statements", []),
            "bullets": {"Limitations": content.get("limitations", []),
                        "Comparison warning": [content["comparison_warning"]] if content.get("comparison_warning") else []},
            "evidence": content.get("evidence", {}),
            "generated_at": content.get("generated_at"),
            "version": content.get("model_workflow_version"),
            "confidence": content.get("confidence"),
        }
    if kind == "impact_narrative":
        st = [s for sec in content.get("sections", {}).values() for s in sec]
        return {"title": content.get("headline", "Intelligence event"), "summary": "", "statements": st, "bullets": {},
                "evidence": content.get("evidence", {}), "generated_at": content.get("generated_at"),
                "version": content.get("model_workflow_version")}
    # executive brief
    sec = content.get("sections", {})
    bullets = {
        "Top developments": [f"{d['headline']} - So what (interpretation): {d['so_what']}" for d in content.get("top_developments", [])],
        "Upcoming catalysts (next 90 days)": [
            f"{c['title']}: {c['expected_date'] or ' to '.join(c.get('window') or [])} "
            f"[{'sourced date' if c['date_basis'] == 'SOURCED' else 'model-inferred window'}]" for c in sec.get("catalysts", [])],
        "New entrants": [e["title"] for e in sec.get("new_entrants", [])],
        "Regulatory / scientific developments": [e["title"] for e in sec.get("regulatory_scientific", [])],
        "Watch items": content.get("watch_items", []),
        "Low-materiality summary": [f"{k}: {v}" for k, v in (sec.get("low_materiality_summary", {}).get("by_type") or {}).items()],
    }
    return {"title": content.get("title", "Executive brief"), "summary": content.get("executive_summary", ""),
            "statements": [], "bullets": bullets, "evidence": {}, "generated_at": content.get("generated_at"),
            "version": content.get("model_workflow_version"), "disclaimer": content.get("disclaimer")}


def to_markdown(kind: str, content: dict[str, Any]) -> str:
    m = _doc_model(kind, content)
    out = [f"# {m['title']}", "", f"_Generated {m['generated_at']} - {m['version']}_", ""]
    if m.get("summary"):
        out += [m["summary"], ""]
    for s in m["statements"]:
        refs = f" [{', '.join(s.get('evidence_ids') or [])}]" if s.get("evidence_ids") else ""
        out.append(f"- **{TYPE_LABEL.get(s['statement_type'], s['statement_type'])}** ({s.get('confidence')}): {s['statement']}{refs}")
    for head, items in m["bullets"].items():
        if items:
            out += ["", f"## {head}"] + [f"- {i}" for i in items]
    if m["evidence"]:
        out += ["", "## Sources"]
        for eid, e in m["evidence"].items():
            out.append(f"- [{eid}] {e.get('source_type')}: {e.get('title') or ''} {e.get('uri') or ''} "
                       f"(retrieved {e.get('retrieved_at')})")
    if m.get("disclaimer"):
        out += ["", f"_{m['disclaimer']}_"]
    return "\n".join(out) + "\n"


def to_docx(kind: str, content: dict[str, Any]) -> bytes:
    from docx import Document

    m = _doc_model(kind, content)
    doc = Document()
    doc.add_heading(m["title"], level=1)
    doc.add_paragraph(f"Generated {m['generated_at']} | {m['version']}").italic = True
    if m.get("summary"):
        doc.add_paragraph(m["summary"])
    for s in m["statements"]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{TYPE_LABEL.get(s['statement_type'], s['statement_type'])} ({s.get('confidence')}): ").bold = True
        p.add_run(s["statement"])
        if s.get("evidence_ids"):
            p.add_run(f" [{', '.join(s['evidence_ids'])}]")
    for head, items in m["bullets"].items():
        if items:
            doc.add_heading(head, level=2)
            for i in items:
                doc.add_paragraph(i, style="List Bullet")
    if m["evidence"]:
        doc.add_heading("Sources", level=2)
        for eid, e in m["evidence"].items():
            doc.add_paragraph(f"[{eid}] {e.get('source_type')}: {e.get('title') or ''} - {e.get('uri') or ''} "
                              f"(retrieved {e.get('retrieved_at')})", style="List Number")
    if m.get("disclaimer"):
        doc.add_paragraph(m["disclaimer"]).italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _latin1(s: str) -> str:
    return (s or "").replace("→", "->").replace("—", "-").replace("–", "-").encode(
        "latin-1", "replace").decode("latin-1")


def _mc(pdf: Any, h: float, text: str) -> None:
    pdf.multi_cell(0, h, text, new_x="LMARGIN", new_y="NEXT", wrapmode="CHAR")


def to_pdf(kind: str, content: dict[str, Any]) -> bytes:
    from fpdf import FPDF

    m = _doc_model(kind, content)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    _mc(pdf, 8, _latin1(m["title"]))
    pdf.set_font("Helvetica", "I", 8)
    _mc(pdf, 5, _latin1(f"Generated {m['generated_at']} | {m['version']}"))
    pdf.set_font("Helvetica", "", 10)
    if m.get("summary"):
        pdf.ln(2)
        _mc(pdf, 5, _latin1(m["summary"]))
    for s in m["statements"]:
        refs = f" [{', '.join(s.get('evidence_ids') or [])}]" if s.get("evidence_ids") else ""
        _mc(pdf, 5, _latin1(f"- {TYPE_LABEL.get(s['statement_type'])} ({s.get('confidence')}): {s['statement']}{refs}"))
    for head, items in m["bullets"].items():
        if items:
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            _mc(pdf, 6, _latin1(head))
            pdf.set_font("Helvetica", "", 10)
            for i in items:
                _mc(pdf, 5, _latin1(f"- {i}"))
    if m["evidence"]:
        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 11)
        _mc(pdf, 6, "Sources")
        pdf.set_font("Helvetica", "", 8)
        for eid, e in m["evidence"].items():
            _mc(pdf, 4, _latin1(f"[{eid}] {e.get('source_type')}: {e.get('title') or ''} {e.get('uri') or ''}"))
    if m.get("disclaimer"):
        pdf.ln(2)
        pdf.set_font("Helvetica", "I", 8)
        _mc(pdf, 4, _latin1(m["disclaimer"]))
    return bytes(pdf.output())


def to_pptx(kind: str, content: dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt

    m = _doc_model(kind, content)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = m["title"][:200]
    s.placeholders[1].text = f"Generated {m['generated_at']}\n{m['version']}"

    def bullet_slide(title: str, lines: list[str]) -> None:
        for i in range(0, len(lines), 6):
            sl = prs.slides.add_slide(prs.slide_layouts[1])
            sl.shapes.title.text = title if i == 0 else f"{title} (cont.)"
            tf = sl.placeholders[1].text_frame
            tf.clear()
            for j, line in enumerate(lines[i:i + 6]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = line[:400]
                p.font.size = Pt(14)

    if m.get("summary"):
        bullet_slide("Summary", [m["summary"]])
    if m["statements"]:
        bullet_slide("Findings", [f"{TYPE_LABEL.get(x['statement_type'])}: {x['statement']}"
                                  + (f" [{', '.join(x['evidence_ids'])}]" if x.get("evidence_ids") else "")
                                  for x in m["statements"]])
    for head, items in m["bullets"].items():
        if items:
            bullet_slide(head, items)
    if m["evidence"]:
        bullet_slide("Sources", [f"[{k}] {v.get('source_type')}: {v.get('uri') or v.get('title')}" for k, v in m["evidence"].items()])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


EXPORTERS = {
    "md": (to_markdown, "text/markdown"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (to_pdf, "application/pdf"),
    "pptx": (to_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
}
